"""docs/08 batch 6.2: the ten weekly numbers. The exit criterion is that
"every tile reconciles exactly with its drill-down (asserted in tests)" —
`test_every_tile_reconciles_with_its_drill_down` is that assertion."""

from __future__ import annotations

from datetime import date, datetime, timedelta

import pytest
from django.utils import timezone
from freezegun import freeze_time

from analytics import metrics
from billing import services as billing_services
from customers.models import Customer
from fulfilment.models import Job, JobAttempt, RouteDay
from growth import services as growth_services
from growth.models import Feedback, PartnerKind
from ordering.models import Order, OrderLine, OrderStatus

pytestmark = pytest.mark.django_db

# A fixed midday-IST instant: "today" must not move under a test, and a
# run that straddles IST midnight would otherwise compare windows built on
# two different days.
NOW = "2026-10-07 06:30:00+00:00"
TODAY = date(2026, 10, 7)


@pytest.fixture(autouse=True)
def frozen_clock():
    with freeze_time(NOW):
        yield


WEEK, WEEK_END = metrics.week_bounds(TODAY - timedelta(days=7))  # last full week
TZ = timezone.get_current_timezone()


def at(day: date, hour: int = 12):
    return timezone.make_aware(datetime.combine(day, datetime.min.time()), TZ) + timedelta(
        hours=hour
    )


def _delivered(hub, customer, service, garment_type, *, day, qty=2, price=1500, code="", apt=None):
    order = Order.objects.create(
        hub=hub,
        customer=customer,
        service=service,
        status=OrderStatus.SCHEDULED,
        referral_code=code,
        apartment=apt,
        declared_total_qty=qty,
        verified_total_qty=qty,
        subtotal_minor=qty * price,
        total_minor=qty * price,
    )
    OrderLine.objects.create(
        hub=hub,
        order=order,
        garment_type=garment_type,
        declared_qty=qty,
        verified_qty=qty,
        unit_price_minor=price,
        line_total_minor=qty * price,
    )
    growth_services.capture_attribution(order)
    order.status = OrderStatus.DELIVERED
    order.delivered_at = at(day)
    order.save(update_fields=["status", "delivered_at"])
    billing_services.issue_invoice(order)
    return order


@pytest.fixture
def week(hub, cluster, apartment, service, garment_type, active_price_list, admin_user):
    partner = growth_services.onboard_partner(
        hub=hub, kind=PartnerKind.WATCHMAN, name="Ramesh", phone="9876500021"
    )
    code = growth_services.create_referral_code(hub=hub, owner_partner=partner).code
    c = {
        n: Customer.objects.create(hub=hub, name=f"{n} Rao", phone=f"+91988800{i:04d}")
        for i, n in enumerate(["Asha", "Bina", "Chitra", "Deepa"])
    }
    # Asha: came via the watchman, first order this week, second order this week too.
    a1 = _delivered(hub, c["Asha"], service, garment_type, day=WEEK, code=code, apt=apartment)
    _delivered(hub, c["Asha"], service, garment_type, day=WEEK + timedelta(days=2), apt=apartment)
    # Bina: an old customer who came back this week.
    _delivered(hub, c["Bina"], service, garment_type, day=WEEK - timedelta(days=20))
    b2 = _delivered(hub, c["Bina"], service, garment_type, day=WEEK + timedelta(days=3), qty=4)
    # Chitra: new this week, one order.
    _delivered(hub, c["Chitra"], service, garment_type, day=WEEK + timedelta(days=4))
    # Deepa: booked, never delivered — a lead, not a customer.
    Order.objects.create(
        hub=hub, customer=c["Deepa"], service=service, status=OrderStatus.CANCELLED
    )

    billing_services.record_order_cost(a1, kind="CONSUMABLE", amount_minor=200)
    billing_services.record_order_cost(b2, kind="LABOUR", amount_minor=900)

    day = RouteDay.objects.create(hub=hub, cluster=cluster, date=WEEK)
    Job.objects.create(
        hub=hub,
        route_day=day,
        order=a1,
        kind="DELIVERY",
        status="DONE",
        slot_end=at(WEEK, 13),
        completed_at=at(WEEK, 12),
    )
    Job.objects.create(
        hub=hub,
        route_day=day,
        order=b2,
        kind="DELIVERY",
        status="DONE",
        slot_end=at(WEEK, 10),
        completed_at=at(WEEK, 12),
    )
    failed = Job.objects.create(
        hub=hub, route_day=day, order=b2, kind="PICKUP", status="FAILED", slot_end=at(WEEK, 9)
    )
    JobAttempt.objects.create(
        job=failed, attempt_no=1, outcome="FAILED", failure_reason="CUSTOMER_ABSENT", at=at(WEEK, 9)
    )

    Feedback.objects.create(hub=hub, order=a1, customer=c["Asha"], rating=5)
    fb = Feedback.objects.create(hub=hub, order=b2, customer=c["Bina"], rating=2)
    Feedback.objects.filter(pk=fb.pk).update(created_at=at(WEEK + timedelta(days=3)))
    Feedback.objects.filter(order=a1).update(created_at=at(WEEK + timedelta(days=1)))
    return c


def test_headline_figures(hub, week):
    def value(key):
        return metrics.METRICS[key][1](hub, WEEK, WEEK_END)

    assert value("new_customers")["value"] == 2  # Asha, Chitra — not Deepa
    repeat = value("repeat_customers")
    assert repeat["value"] == 1  # Bina came back
    assert repeat["cohort_size"] == 2 and repeat["cohort_repeated"] == 1  # Asha reordered
    opc = value("orders_per_customer")
    assert (opc["orders"], opc["active_customers"]) == (4, 3)
    assert value("referrals")["value"] == 1
    assert value("referrals")["leaderboard"] == [{"partner": "Ramesh", "new_customers": 1}]
    aov = value("average_order_value")
    assert aov["value"] == round((3000 + 3000 + 6000 + 3000) / 4)
    contribution = value("contribution")
    assert contribution["margin_minor"] == 15000 - 200 - 900
    on_time = value("on_time")
    assert on_time["value"] == round(100 / 3, 1)  # 1 on time, 1 late, 1 failed
    assert on_time["excluding_customer_caused"] == 50.0
    fb = value("feedback")
    assert fb["value"] == 3.5 and fb["detractors"] == 1 and fb["response_rate"] == 50.0


def test_every_tile_reconciles_with_its_drill_down(hub, week):
    for key, (_label, fn, _) in metrics.METRICS.items():
        result = fn(hub, WEEK, WEEK_END)
        rows = result["rows"]
        if key in ("new_customers", "repeat_customers", "referrals"):
            assert result["value"] == len(rows), key
        elif key == "orders_per_customer":
            assert result["orders"] == sum(r["orders"] for r in rows)
            assert result["active_customers"] == len(rows)
        elif key == "acquisition_cost":
            assert result["new_customers"] == sum(r["new_customers"] for r in rows)
        elif key == "apartments":
            assert (
                sum(r["orders"] for r in rows)
                == metrics.delivered_orders(hub, WEEK, WEEK_END).count()
            )
        elif key == "average_order_value":
            assert result["revenue_minor"] == sum(r["net_minor"] for r in rows)
            assert result["orders"] == len(rows)
        elif key == "contribution":
            assert result["margin_minor"] == sum(r["margin_minor"] for r in rows)
            for kind in metrics.COST_KINDS:
                assert result["costs_minor"][kind] == sum(r[kind] for r in rows)
        elif key == "on_time":
            on_time = sum(r["result"] == "ON_TIME" for r in rows)
            assert result["value"] == round(100 * on_time / len(rows), 1)
        elif key == "feedback":
            assert result["responses"] == len(rows)
            assert result["value"] == round(sum(r["rating"] for r in rows) / len(rows), 2)
        else:
            pytest.fail(f"no reconciliation for {key}")


def test_weekly_pack_has_trend_and_restricts_money_for_admin(hub, week):
    founder = metrics.weekly(hub, WEEK, include_money=True)
    tiles = {t["key"]: t for t in founder["tiles"]}
    assert list(tiles) == list(metrics.METRICS)
    assert len(tiles["new_customers"]["trend"]) == metrics.SPARK_WEEKS
    assert tiles["new_customers"]["trend"][-1]["value"] == 2
    assert "rows" not in tiles["new_customers"]

    admin = {t["key"]: t for t in metrics.weekly(hub, WEEK, include_money=False)["tiles"]}
    assert admin["contribution"] == {
        "key": "contribution",
        "label": "Money made per order",
        "restricted": True,
    }
    assert admin["acquisition_cost"]["restricted"]


def test_empty_week_is_empty_not_an_error(hub):
    pack = metrics.weekly(hub, WEEK, include_money=True)
    tiles = {t["key"]: t for t in pack["tiles"]}
    assert tiles["new_customers"]["value"] == 0
    assert tiles["average_order_value"]["value"] is None
    assert tiles["on_time"]["value"] is None


# ------------------------------------------------------------------- API


def test_weekly_api_roles(api_client, founder_user, admin_user, operator_user, hub, week):
    url = f"/api/v1/analytics/weekly?hub={hub.id}&week={WEEK + timedelta(days=3)}"
    api_client.force_authenticate(user=operator_user)
    assert api_client.get(url).status_code == 403

    api_client.force_authenticate(user=admin_user)
    resp = api_client.get(url)
    assert resp.status_code == 200
    assert resp.data["week_start"] == str(WEEK)
    assert (
        api_client.get(
            f"/api/v1/analytics/weekly/contribution/rows?hub={hub.id}&week={WEEK}"
        ).status_code
        == 403
    )
    rows = api_client.get(f"/api/v1/analytics/weekly/new_customers/rows?hub={hub.id}&week={WEEK}")
    assert rows.status_code == 200 and rows.data["value"] == len(rows.data["rows"]) == 2

    api_client.force_authenticate(user=founder_user)
    assert (
        api_client.get(
            f"/api/v1/analytics/weekly/contribution/rows?hub={hub.id}&week={WEEK}"
        ).status_code
        == 200
    )
    assert api_client.get(f"/api/v1/analytics/weekly/nope/rows?week={WEEK}").status_code == 404


def test_exports(api_client, founder_user, hub, week):
    import io

    from openpyxl import load_workbook

    api_client.force_authenticate(user=founder_user)
    xlsx = api_client.get(f"/api/v1/analytics/weekly/export.xlsx?hub={hub.id}&week={WEEK}")
    assert xlsx.status_code == 200
    wb = load_workbook(io.BytesIO(xlsx.content))
    assert wb.sheetnames[0] == "Summary" and wb.sheetnames[-1] == "Parameters"
    assert "Money made per order" in wb.sheetnames
    aov_row = next(
        r for r in wb["Summary"].iter_rows(values_only=True) if r[0] == "Average order value"
    )
    assert aov_row[1] == 37.5  # rupees as a number, not "₹37.50"

    csv = api_client.get(
        f"/api/v1/analytics/weekly/new_customers/export.csv?hub={hub.id}&week={WEEK}"
    )
    assert csv.status_code == 200
    assert csv.content.decode().splitlines()[0].startswith("customer,name,phone")
